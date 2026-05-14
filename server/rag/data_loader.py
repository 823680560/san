"""
数据加载器 - 加载多种格式的数据文件
支持: Excel (.xlsx/.xls), JSON, CSV, PDF, TXT, Word (.doc/.docx)
"""
import json
import os
import pandas as pd
from typing import List
from llama_index.core import Document


class DataLoader:
    """数据加载器，负责加载和预处理各类数据源"""

    @staticmethod
    def load_excel_data(file_path: str) -> List[Document]:
        """
        加载Excel招标项目数据
        返回 Document 列表
        """
        df = pd.read_excel(file_path, sheet_name="Sheet1")
        documents = []

        for idx, row in df.iterrows():
            # 将每一行转换为结构化的文本描述
            text_parts = []
            metadata = {}

            for col in df.columns:
                value = row[col]
                if pd.notna(value) and value != "":
                    text_parts.append(f"{col}: {value}")
                    metadata[col] = str(value)

            text = "\n".join(text_parts)
            doc = Document(
                text=text,
                metadata={
                    "source": "excel",
                    "row_index": idx,
                    "title": str(row.get("标题", "")),
                    "category": str(row.get("类别", "")),
                    "project_name": str(row.get("项目名称", "")),
                    "project_id": str(row.get("项目编号", "")),
                    "budget": str(row.get("预算", "")),
                    "region": f"{row.get('省份', '')} {row.get('市区', '')} {row.get('县城', '')}".strip(),
                }
            )
            documents.append(doc)

        print(f"从Excel加载了 {len(documents)} 条招标项目数据")
        return documents

    @staticmethod
    def load_json_data(file_path: str) -> List[Document]:
        """
        加载JSON政策法规数据
        返回 Document 列表
        """
        with open(file_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        documents = []
        seen_titles = set()  # 去重

        for item in data:
            title = item.get("政策标题", "")
            if title in seen_titles:
                continue
            seen_titles.add(title)

            content = item.get("政策法规内容", "")
            publish_time = item.get("发布时间", "")
            effective_time = item.get("施行时间", "")
            source = item.get("来源网站", "")

            text = f"标题: {title}\n"
            if publish_time:
                text += f"发布时间: {publish_time}\n"
            if effective_time:
                text += f"施行时间: {effective_time}\n"
            text += f"来源: {source}\n\n"
            text += content

            doc = Document(
                text=text,
                metadata={
                    "source": "json",
                    "title": title,
                    "publish_time": publish_time,
                    "effective_time": effective_time,
                    "source_website": source,
                }
            )
            documents.append(doc)

        print(f"从JSON加载了 {len(documents)} 条政策法规数据")
        return documents

    @staticmethod
    def load_csv_data(file_path: str) -> List[Document]:
        """
        加载CSV数据文件
        返回 Document 列表
        """
        df = pd.read_csv(file_path)
        documents = []

        for idx, row in df.iterrows():
            text_parts = []
            metadata = {}

            for col in df.columns:
                value = row[col]
                if pd.notna(value) and str(value).strip() != "":
                    text_parts.append(f"{col}: {value}")
                    metadata[col] = str(value)

            text = "\n".join(text_parts)
            doc = Document(
                text=text,
                metadata={
                    "source": "csv",
                    "row_index": idx,
                    "file_name": os.path.basename(file_path),
                }
            )
            documents.append(doc)

        print(f"从CSV加载了 {len(documents)} 条数据")
        return documents

    @staticmethod
    def load_txt_data(file_path: str) -> List[Document]:
        """
        加载TXT文本文件
        返回 Document 列表
        """
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()

        doc = Document(
            text=text,
            metadata={
                "source": "txt",
                "file_name": os.path.basename(file_path),
                "title": os.path.splitext(os.path.basename(file_path))[0],
            }
        )
        print(f"从TXT加载了 1 个文档 ({len(text)} 字符)")
        return [doc]

    @staticmethod
    def load_pdf_data(file_path: str) -> List[Document]:
        """
        加载PDF文件
        使用 PyMuPDF (fitz) 解析，每页作为一个 Document
        返回 Document 列表
        """
        try:
            import fitz  # PyMuPDF
        except ImportError:
            raise ImportError("请安装 PyMuPDF: pip install PyMuPDF")

        documents = []
        doc = fitz.open(file_path)

        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text()

            if text.strip():
                doc_item = Document(
                    text=text,
                    metadata={
                        "source": "pdf",
                        "file_name": os.path.basename(file_path),
                        "title": os.path.splitext(os.path.basename(file_path))[0],
                        "page_number": page_num + 1,
                        "total_pages": len(doc),
                    }
                )
                documents.append(doc_item)

        doc.close()
        print(f"从PDF加载了 {len(documents)} 页内容")
        return documents

    @staticmethod
    def load_law_pdf_data(file_path: str) -> List[Document]:
        """
        加载法规类 PDF 文件，提取元数据对齐 JSON 法规格式

        与通用 load_pdf_data 的区别：
        - 从 PDF 内置元数据或文件名中提取 title/publish_time
        - 元数据结构对齐 load_json_data，便于合并到 law 知识库
        返回 Document 列表
        """
        import re

        # ① 用已有方法加载 PDF 页面
        page_docs = DataLoader.load_pdf_data(file_path)
        if not page_docs:
            return []

        # ② 从 PDF 内置信息提取元数据
        try:
            import fitz
            pdf_doc = fitz.open(file_path)
            pdf_meta = pdf_doc.metadata or {}
            pdf_doc.close()
        except Exception:
            pdf_meta = {}

        # 书名：优先用 PDF 内置 title，否则从文件名清洗
        raw_title = pdf_meta.get("title", "")
        if not raw_title:
            raw_title = os.path.splitext(os.path.basename(file_path))[0]
        # 清洗：去掉出版社、Z-Library、OCR 标记等括号内附注
        clean_title = re.sub(r"\s*[\(（].*?[\)）]\s*", "", raw_title)
        clean_title = re.sub(r"\s+", "", clean_title).strip()
        if not clean_title:
            clean_title = raw_title.strip()

        # 发布时间：从 PDF creationDate 提取年份
        publish_time = ""
        creation_date = pdf_meta.get("creationDate", "")
        if creation_date:
            m = re.search(r"(\d{4})", creation_date)
            if m:
                publish_time = m.group(1)

        # ③ 为每个页面重新设置元数据（同时清洗 OCR 产生的非法代理字符）
        enriched_docs = []
        for page_doc in page_docs:
            page_meta = dict(page_doc.metadata)
            sanitized_text = re.sub(r'[\ud800-\udfff]', '', page_doc.text)
            enriched_docs.append(Document(
                text=sanitized_text,
                metadata={
                    "source": "pdf_book",
                    "title": clean_title,
                    "publish_time": publish_time,
                    "effective_time": "",
                    "source_website": "PDF导入",
                    "file_name": os.path.basename(file_path),
                    "page_number": page_meta.get("page_number", 0),
                    "total_pages": page_meta.get("total_pages", 0),
                }
            ))

        print(f"从法规PDF加载了 {len(enriched_docs)} 页内容（书名: {clean_title}）")
        return enriched_docs

    @staticmethod
    def load_docx_data(file_path: str) -> List[Document]:
        """
        加载 .docx Word 文档
        使用 python-docx 解析
        返回 Document 列表
        """
        try:
            from docx import Document as DocxDocument
        except ImportError:
            raise ImportError("请安装 python-docx: pip install python-docx")

        docx_doc = DocxDocument(file_path)

        # 提取所有段落文本
        paragraphs = [p.text for p in docx_doc.paragraphs if p.text.strip()]
        full_text = "\n".join(paragraphs)

        # 提取表格内容
        tables_text = []
        for table in docx_doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells])
                if row_text.strip():
                    tables_text.append(row_text)

        if tables_text:
            full_text += "\n\n【表格内容】\n" + "\n".join(tables_text)

        doc = Document(
            text=full_text,
            metadata={
                "source": "docx",
                "file_name": os.path.basename(file_path),
                "title": os.path.splitext(os.path.basename(file_path))[0],
            }
        )
        print(f"从DOCX加载了 1 个文档 ({len(full_text)} 字符)")
        return [doc]

    @staticmethod
    def load_doc_data(file_path: str) -> List[Document]:
        """
        加载 .doc 旧版 Word 文档
        优先尝试用 python-docx 读取，失败则用 textract
        返回 Document 列表
        """
        try:
            # 先尝试用 python-docx 读取（部分 .doc 可读）
            return DataLoader.load_docx_data(file_path)
        except Exception:
            try:
                import textract
                text = textract.process(file_path).decode('utf-8')
                doc = Document(
                    text=text,
                    metadata={
                        "source": "doc",
                        "file_name": os.path.basename(file_path),
                        "title": os.path.splitext(os.path.basename(file_path))[0],
                    }
                )
                print(f"从DOC加载了 1 个文档 ({len(text)} 字符)")
                return [doc]
            except ImportError:
                raise ImportError("请安装 textract: pip install textract")
            except Exception as e:
                print(f"DOC解析失败: {e}")
                # 最后尝试用 antiword 命令行工具
                import subprocess
                try:
                    result = subprocess.run(
                        ['antiword', file_path],
                        capture_output=True,
                        text=True,
                        encoding='utf-8'
                    )
                    if result.returncode == 0:
                        text = result.stdout
                        doc = Document(
                            text=text,
                            metadata={
                                "source": "doc",
                                "file_name": os.path.basename(file_path),
                                "title": os.path.splitext(os.path.basename(file_path))[0],
                            }
                        )
                        print(f"从DOC(antiword)加载了 1 个文档 ({len(text)} 字符)")
                        return [doc]
                except FileNotFoundError:
                    pass
                raise ValueError(f"无法解析DOC文件: {file_path}")

    @staticmethod
    def load_file(file_path: str) -> List[Document]:
        """
        根据文件扩展名自动选择加载方式
        统一的入口函数

        Args:
            file_path: 文件路径

        Returns:
            Document 列表
        """
        ext = os.path.splitext(file_path)[1].lower()

        loaders = {
            '.xlsx': DataLoader.load_excel_data,
            '.xls': DataLoader.load_excel_data,
            '.json': DataLoader.load_json_data,
            '.csv': DataLoader.load_csv_data,
            '.txt': DataLoader.load_txt_data,
            '.pdf': DataLoader.load_pdf_data,
            '.docx': DataLoader.load_docx_data,
            '.doc': DataLoader.load_doc_data,
        }

        loader = loaders.get(ext)
        if loader is None:
            raise ValueError(f"不支持的文件格式: {ext}，支持的格式: {list(loaders.keys())}")

        return loader(file_path)
       